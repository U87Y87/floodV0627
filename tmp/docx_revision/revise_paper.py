from copy import deepcopy
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

SOURCE = Path(r'C:\Users\11823\Desktop\小论文\论文讨论0903.docx')
OUT = Path(r'D:\github\V0622\output\documents\论文讨论0903_原型系统研发修订.docx')
doc = Document(SOURCE)

def replace_paragraph(paragraph, text):
    formatting = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.clear()
    run = paragraph.add_run(text)
    if formatting is not None:
        run._r.insert(0, formatting)

paragraph = next(p for p in doc.paragraphs if p.text.startswith('原型系统研发的软硬件配置'))
replacement = (
    '原型系统采用浏览器/服务器（B/S）架构，在Windows 11环境下完成开发与测试，软硬件配置见表2。'
    '系统由客户端、服务端及数据处理与模拟服务组成，各组件通过HTTP接口进行数据交互。'
    '服务端以Node.js为运行环境，利用Vite提供开发服务、接口处理与请求转发；'
    '数据处理模块基于Python实现，结合Flask、Rasterio、NumPy和pyproj完成DEM格式转换、地表糙率计算、坐标处理及格网转换等操作，'
    '洪水数值计算通过HTTP接口调用独立模拟服务完成。'
    '客户端采用Vue 3构建交互界面，利用Cesium.js组织三维地形、遥感影像和洪水三角网格，'
    '实现模拟过程逐帧播放、水深与淹没面积动态统计，以及建筑和道路累计受淹状态展示。'
    '系统通过JSON及TIF、ASC、DAT、TXT等文件交换参数与计算结果；'
    '任务交互模块采用关键词规则匹配，为数据处理、洪水模拟和可视化等步骤提供功能推荐与操作引导。'
)
replace_paragraph(paragraph, replacement)

caption = next(p for p in doc.paragraphs if p.text.startswith('表2') and '配置' in p.text)
replace_paragraph(caption, '表2 原型系统实验环境与软件配置')
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.keep_with_next = True
caption.paragraph_format.space_before = Pt(8)
caption.paragraph_format.space_after = Pt(5)

following = next(p for p in doc.paragraphs if p.text.startswith('基于上述软硬件条件'))
replace_paragraph(following,
    '基于上述软硬件环境，搭建山洪模拟与三维可视化原型系统，系统Web端界面如图8所示。'
    '系统主要包括数据处理、洪水模拟、格网转换、三维可视化、案例展示和任务交互模块，'
    '支持从模拟数据准备、模型调用到结果展示的流程衔接。')

rows = [
    ('硬件与操作系统', 'CPU', 'Intel Core i9-14900HX'),
    ('硬件与操作系统', '内存', '32.0 GB'),
    ('硬件与操作系统', '操作系统', 'Windows 11'),
    ('服务端软件', 'JavaScript运行环境', 'Node.js 23.10.0'),
    ('服务端软件', 'Python运行环境', 'Python 3.13.9'),
    ('服务端软件', '开发服务与请求转发', 'Vite 8.0.16'),
    ('服务端软件', '数据处理框架与库', 'Flask、Rasterio、NumPy、pyproj'),
    ('服务端软件', '洪水模拟服务', '通过HTTP接口调用独立数值模拟服务'),
    ('客户端软件', '前端框架', 'Vue 3.5.34'),
    ('客户端软件', '三维可视化引擎', 'Cesium.js 1.142.0'),
    ('客户端软件', '运行方式', '支持WebGL的现代浏览器'),
    ('网络与数据交互', '通信协议', 'TCP/IP、HTTP'),
    ('网络与数据交互', '接口数据格式', 'JSON'),
    ('网络与数据交互', '数据文件格式', 'TIF、ASC、DAT、TXT、JSON'),
    ('网络与数据交互', '实验访问方式', '本机部署，通过浏览器访问'),
    ('任务交互', '功能引导方式', '基于关键词规则匹配推荐系统工具'),
]
old_table = next(t for t in doc.tables if any('Nginx' in c.text for r in t.rows for c in r.cells))
table = doc.add_table(rows=len(rows) + 1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Mm(30), Mm(39), Mm(77)]
for col, width in zip(table.columns, widths):
    col.width = width
for i, values in enumerate([('类别', '项目名称', '具体信息'), *rows]):
    row = table.rows[i]
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    tr_pr.append(cant_split)
    if i == 0:
        repeat = OxmlElement('w:tblHeader')
        tr_pr.append(repeat)
    for j, text in enumerate(values):
        cell = row.cells[j]
        cell.width = widths[j]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = text
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or j == 0 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.paragraph_format.keep_with_next = i < len(rows)
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            run.font.bold = i == 0
            run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for edge, value in [('top', 70), ('bottom', 70), ('left', 70), ('right', 70)]:
            node = OxmlElement(f'w:{edge}')
            node.set(qn('w:w'), str(value))
            node.set(qn('w:type'), 'dxa')
            margins.append(node)
        tc_pr.append(margins)

# Merge consecutive category labels, without inheriting the incorrect source grouping.
start = 1
for i in range(2, len(rows) + 2):
    if i == len(rows) + 1 or rows[i - 1][0] != rows[start - 1][0]:
        if i - start > 1:
            merged = table.cell(start, 0).merge(table.cell(i - 1, 0))
            merged.text = rows[start - 1][0]
            merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = merged.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                run._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        start = i

# Retain the paper's three-line table convention.
tbl_pr = table._tbl.tblPr
borders = OxmlElement('w:tblBorders')
for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
    node = OxmlElement(f'w:{edge}')
    node.set(qn('w:val'), 'single' if edge in ['top', 'bottom'] else 'nil')
    node.set(qn('w:sz'), '10')
    node.set(qn('w:color'), '000000')
    borders.append(node)
tbl_pr.append(borders)
for cell in table.rows[0].cells:
    cell_borders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:color'), '000000')
    cell_borders.append(bottom)
    cell._tc.get_or_add_tcPr().append(cell_borders)
old_table._tbl.addprevious(table._tbl)
old_table._tbl.getparent().remove(old_table._tbl)

# Only replace document.xml; all equations, images, relationships and other
# original package parts remain byte-for-byte unchanged.
buffer = BytesIO()
doc.save(buffer)
buffer.seek(0)
with ZipFile(buffer) as edited:
    document_xml = edited.read('word/document.xml')
OUT.parent.mkdir(parents=True, exist_ok=True)
with ZipFile(SOURCE) as original, ZipFile(OUT, 'w') as result:
    for item in original.infolist():
        result.writestr(item, document_xml if item.filename == 'word/document.xml' else original.read(item.filename))

verified = Document(OUT)
assert len(verified.tables) == len(doc.tables)
assert '见表2' in next(p.text for p in verified.paragraphs if p.text.startswith('原型系统采用'))
assert not any(word in '\n'.join(c.text for r in verified.tables[-1].rows for c in r.cells) for word in ['Nginx', 'AutoGen', 'Neo4j', 'Token', 'Ubuntu'])
print(OUT)
